
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def export_comparison_dataset(history):
    """Export algorithm run history to an Excel .xlsx file."""

    if not history:
        return None

    # Collect all columns recorded in the history
    fieldnames = []
    for record in history:
        for key in record:
            if key not in fieldnames:
                fieldnames.append(key)

    # Create workbook and worksheet
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Comparison Dataset"

    # Write column headers
    worksheet.append(fieldnames)

    # Write data rows
    for record in history:
        worksheet.append([
            record.get(field, "")
            for field in fieldnames
        ])

    # Format the header row
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            "solid",
            fgColor="4472C4"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    # Adjust column widths
    for column_index, column_cells in enumerate(
        worksheet.columns, start=1
    ):
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = min(max_length + 2, 35)

    # Freeze the header row and enable filtering
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    # Save the Excel file beside main.py
    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S_%f"
    )
    filename = f"comparison_dataset_{timestamp}.xlsx"
    output_path = Path(__file__).resolve().parent / filename

    workbook.save(output_path)
    return output_path
